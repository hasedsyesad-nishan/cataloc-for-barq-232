#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy
import os

IMAGES_DIR = "/tmp/docx_extracted/word/media"

def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'right')
    pPr.append(jc)

def set_font(run, name='Tajawal', size=None, bold=False, color=None):
    run.font.name = name
    run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    run._r.rPr.append(rFonts)
    if size:
        run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    if color:
        run.font.color.rgb = RGBColor(*color)

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if val:
            el = OxmlElement(f'w:{side}')
            el.set(qn('w:val'), val.get('val', 'single'))
            el.set(qn('w:sz'), str(val.get('sz', 4)))
            el.set(qn('w:color'), val.get('color', '000000'))
            tcBorders.append(el)
    tcPr.append(tcBorders)

def add_heading(doc, text, level=1, align='right', color=(0x1A, 0x56, 0x9A)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == 'right' else WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    run = p.add_run(text)
    sizes = {1: 26, 2: 20, 3: 16, 4: 14}
    set_font(run, size=sizes.get(level, 14), bold=True, color=color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_body(doc, text, size=12, color=(0x33, 0x33, 0x33), align='right', bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == 'right' else WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    run = p.add_run(f'• {text}')
    set_font(run, size=size, color=(0x33, 0x33, 0x33))
    p.paragraph_format.space_after = Pt(2)
    return p

def add_separator(doc, color='1A569A'):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_image_centered(doc, img_path, width=Cm(10)):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=width)
        p.paragraph_format.space_after = Pt(8)
    return

def add_info_box(doc, title, content, bg='E8F0FA', border_color='1A569A'):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    cell = table.cell(0, 0)
    set_cell_bg(cell, bg)
    set_cell_borders(cell,
        top={'val': 'single', 'sz': 8, 'color': border_color},
        bottom={'val': 'single', 'sz': 8, 'color': border_color},
        left={'val': 'single', 'sz': 8, 'color': border_color},
        right={'val': 'single', 'sz': 8, 'color': border_color},
    )
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p1)
    r1 = p1.add_run(title)
    set_font(r1, size=13, bold=True, color=(0x1A, 0x56, 0x9A))

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p2)
    r2 = p2.add_run(content)
    set_font(r2, size=11, color=(0x33, 0x33, 0x33))

    for par in cell.paragraphs:
        par.paragraph_format.space_before = Pt(4)
        par.paragraph_format.space_after = Pt(4)

    table.style = 'Table Grid'
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_step_box(doc, number, title, content):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.columns[0].width = Cm(1.8)
    table.columns[1].width = Cm(13)

    # Number cell
    c0 = table.cell(0, 0)
    set_cell_bg(c0, '1A569A')
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(number))
    set_font(r, size=20, bold=True, color=(0xFF, 0xFF, 0xFF))
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Content cell
    c1 = table.cell(0, 1)
    set_cell_bg(c1, 'F0F5FB')
    p2 = c1.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p2)
    r2 = p2.add_run(title)
    set_font(r2, size=13, bold=True, color=(0x1A, 0x56, 0x9A))

    p3 = c1.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p3)
    r3 = p3.add_run(content)
    set_font(r3, size=11, color=(0x33, 0x33, 0x33))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_specs_table(doc, title, rows_data):
    add_heading(doc, title, level=3)
    table = doc.add_table(rows=len(rows_data)+1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.style = 'Table Grid'

    # Header
    h0 = table.cell(0, 0)
    h1 = table.cell(0, 1)
    set_cell_bg(h0, '1A569A')
    set_cell_bg(h1, '1A569A')
    for cell, text in [(h0, 'المواصفة'), (h1, 'القيمة')]:
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=12, bold=True, color=(0xFF, 0xFF, 0xFF))

    for i, (spec, val) in enumerate(rows_data):
        row = table.rows[i+1]
        bg = 'F7FAFF' if i % 2 == 0 else 'FFFFFF'
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_bg(c0, bg)
        set_cell_bg(c1, bg)

        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p0)
        r0 = p0.add_run(spec)
        set_font(r0, size=11, bold=True, color=(0x22, 0x22, 0x22))

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_rtl(p1)
        r1 = p1.add_run(val)
        set_font(r1, size=11, color=(0x44, 0x44, 0x44))

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(docx_breaks.WD_BREAK.PAGE)

from docx.enum.text import WD_BREAK
import docx.enum.text as docx_breaks

# ==============================
# BUILD DOCUMENT
# ==============================
doc = Document()

# Page setup - A4
from docx.oxml import OxmlElement as OE
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)

# RTL document setting
settings = doc.settings.element
bidi = OxmlElement('w:bidi')
settings.append(bidi)

# ============================================================
# COVER PAGE
# ============================================================
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(80)
r = p.add_run('نظام برق')
set_font(r, size=48, bold=True, color=(0x1A, 0x56, 0x9A))

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('BARQ — Smart Energy Management')
set_font(r2, size=18, color=(0x77, 0x77, 0x77))

add_separator(doc)

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.paragraph_format.space_before = Pt(10)
r3 = p3.add_run('نظام ذكي لمراقبة وإدارة منظومات الطاقة الشمسية')
set_font(r3, size=16, color=(0x33, 0x33, 0x33))

# Cover images
p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_img.paragraph_format.space_before = Pt(20)
run_img = p_img.add_run()
img1 = f"{IMAGES_DIR}/image1.jpeg"
img6 = f"{IMAGES_DIR}/image6.jpeg"
if os.path.exists(img1):
    run_img.add_picture(img1, width=Cm(8))

p_img2 = doc.add_paragraph()
p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_img2 = p_img2.add_run()
if os.path.exists(img6):
    run_img2.add_picture(img6, width=Cm(12))

p_footer = doc.add_paragraph()
p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_footer.paragraph_format.space_before = Pt(30)
r_f = p_footer.add_run('كتالوج الأجهزة ودليل التركيب والإعداد')
set_font(r_f, size=14, bold=True, color=(0x1A, 0x56, 0x9A))

p_footer2 = doc.add_paragraph()
p_footer2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f2 = p_footer2.add_run('شركة برق لحلول الطاقة النظيفة — صناعة سورية')
set_font(r_f2, size=12, color=(0x77, 0x77, 0x77))

# Page break after cover
p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 2: ما هو نظام برق؟
# ============================================================
add_heading(doc, 'ما هو نظام برق؟', level=1, align='right')
add_separator(doc)

add_body(doc,
    'برق هو نظام ذكي لمراقبة وإدارة منظومات الطاقة الشمسية المنزلية. يرتبط مباشرة بجهاز الإنفرتر ويقرأ بياناته الفعلية لحظة بلحظة، مما يتيح لك معرفة حالة بطاريتك ومصدر الطاقة في أي وقت — سواء كنت في المنزل أو على بُعد آلاف الكيلومترات.',
    size=12
)

doc.add_paragraph()

# Flow diagram as table
add_heading(doc, 'كيف يعمل النظام؟', level=2)
flow_table = doc.add_table(rows=1, cols=7)
flow_table.alignment = WD_TABLE_ALIGNMENT.CENTER
flow_items = [
    ('⚡', 'الإنفرتر', 'Voltronic\nGrowatt'),
    ('←', '', ''),
    ('📡', 'المرسلة', 'RS-232 / RJ45'),
    ('←', '', ''),
    ('❄️🔥', 'البراد / السخان', 'ESP-NOW لاسلكي'),
    ('←', '', ''),
    ('📱', 'التطبيق', 'WiFi / إنترنت'),
]
for i, (icon, name, sub) in enumerate(flow_items):
    c = flow_table.cell(0, i)
    if icon == '←':
        set_cell_bg(c, 'FFFFFF')
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('←')
        set_font(r, size=20, bold=True, color=(0x1A, 0x56, 0x9A))
    else:
        set_cell_bg(c, 'E8F0FA')
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_icon = p.add_run(icon + '\n')
        set_font(r_icon, size=18)
        r_name = p.add_run(name + '\n')
        set_font(r_name, size=10, bold=True, color=(0x1A, 0x56, 0x9A))
        r_sub = p.add_run(sub)
        set_font(r_sub, size=8, color=(0x77, 0x77, 0x77))

doc.add_paragraph()

add_heading(doc, 'مميزات النظام', level=2)
features = [
    ('🔍 مراقبة فورية', 'اعرف نسبة شحن بطاريتك وجهدها وطاقة الألواح الشمسية واستهلاك الأحمال — كل ذلك في الوقت الفعلي من أي مكان.'),
    ('🎛️ تحكم ذكي', 'اضبط متى يشتغل البراد أو السخان بناءً على نسبة الشحن، وحدد الأولويات لتحافظ على عمر بطاريتك.'),
    ('📊 تحليل الطاقة', 'مخطط بياني لآخر 24 ساعة يوضح سلوك البطارية، وساعات تشغيل كل جهاز خلال اليوم.'),
    ('🌐 لا يحتاج إنترنت دائماً', 'يعمل النظام محلياً حتى بدون إنترنت. الإنترنت يضيف إمكانية المراقبة عن بُعد فقط.'),
    ('🔋 حماية البطارية', 'إيقاف تلقائي للأجهزة عند انخفاض الشحن لمستوى محدد، مما يطيل عمر البطارية.'),
    ('🇸🇾 صناعة سورية', 'جميع الأجهزة مصنوعة ومطورة محلياً في سوريا مع دعم فني مباشر.'),
]
for title, desc in features:
    add_info_box(doc, title, desc)

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 3: المرسلة
# ============================================================
add_heading(doc, 'أولاً: المرسلة — قلب النظام', level=1)
add_separator(doc)
add_body(doc, 'تتصل مباشرة بالإنفرتر عبر كابل RJ45 على مخرج COMM، وتقرأ بيانات البطارية والطاقة الشمسية، ثم تبثها لاسلكياً للأجهزة وترفعها للإنترنت.', size=12)

doc.add_paragraph()

# Device image
img2 = f"{IMAGES_DIR}/image2.png"
if os.path.exists(img2):
    p_i = doc.add_paragraph()
    p_i.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_i = p_i.add_run()
    run_i.add_picture(img2, width=Cm(9))
    p_c = doc.add_paragraph()
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run('جهاز المرسلة BARQ')
    set_font(r_c, size=10, color=(0x77, 0x77, 0x77))

doc.add_paragraph()

add_heading(doc, 'مواصفات المرسلة', level=2)
specs_mursila = [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال بالإنفرتر', 'RS-232 / RJ45 (مخرج COMM)'),
    ('الاتصال اللاسلكي', 'ESP-NOW — مدى يصل 200 متر'),
    ('الاتصال بالإنترنت', 'WiFi 2.4GHz'),
    ('التغذية', 'DC 12V'),
    ('دورة رفع البيانات', 'كل دقيقتين تلقائياً'),
    ('الأنفرترات المدعومة', 'Voltronic / Growatt وما يدعم RS-232'),
]
add_specs_table(doc, '', specs_mursila)

add_heading(doc, 'وظائف زر الضبط', level=2)
btn_table = doc.add_table(rows=4, cols=2)
btn_table.style = 'Table Grid'
btn_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
btn_data = [
    ('الزر وطريقة الضغط', 'الوظيفة'),
    ('ضغطة واحدة', 'وضع WiFi — يبث شبكة BARQ للضبط'),
    ('ضغطتان متتاليتان', 'وضع الاقتران مع المستقبلات'),
    ('ضغط مستمر 10 ثوانٍ', 'إعادة الضبط للإعدادات المصنعية'),
]
for i, (col1, col2) in enumerate(btn_data):
    row = btn_table.rows[i]
    c0, c1 = row.cells[0], row.cells[1]
    if i == 0:
        set_cell_bg(c0, '1A569A')
        set_cell_bg(c1, '1A569A')
        for c, t in [(c0, col1), (c1, col2)]:
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(t)
            set_font(r, size=12, bold=True, color=(0xFF,0xFF,0xFF))
    else:
        bg = 'F7FAFF' if i % 2 == 1 else 'FFFFFF'
        set_cell_bg(c0, bg)
        set_cell_bg(c1, bg)
        for c, t in [(c0, col1), (c1, col2)]:
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)
            r = p.add_run(t)
            set_font(r, size=11, color=(0x33,0x33,0x33))

doc.add_paragraph()

# Extra images: back/side
img3 = f"{IMAGES_DIR}/image3.png"
img4 = f"{IMAGES_DIR}/image4.png"
if os.path.exists(img3) or os.path.exists(img4):
    add_heading(doc, 'صور تفصيلية للجهاز', level=3)
    pics_table = doc.add_table(rows=1, cols=2)
    pics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if os.path.exists(img3):
        c0 = pics_table.cell(0, 0)
        p = c0.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(img3, width=Cm(7))
        p2 = c0.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run('الخلفية: مدخل DC 12V + منفذ RJ45')
        set_font(r2, size=9, color=(0x77,0x77,0x77))
    if os.path.exists(img4):
        c1 = pics_table.cell(0, 1)
        p = c1.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(img4, width=Cm(7))
        p2 = c1.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run('الجانب: زر الضبط الأبيض')
        set_font(r2, size=9, color=(0x77,0x77,0x77))

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 4: البراد والمكيف
# ============================================================
add_heading(doc, 'ثانياً: جهاز البراد / المكيف', level=1)
add_separator(doc)
add_body(doc, 'يتحكم في تشغيل وإيقاف الثلاجة، الفريزر، الكولر، أو المكيف بشكل تلقائي بناءً على نسبة شحن البطارية التي تحددها أنت.', size=12)

doc.add_paragraph()

# Images
img5 = f"{IMAGES_DIR}/image5.png"
img7 = f"{IMAGES_DIR}/image7.png"
if os.path.exists(img5) or os.path.exists(img7):
    pics_table2 = doc.add_table(rows=1, cols=2)
    pics_table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col_idx, (img_path, caption) in enumerate([(img5, 'جهاز البراد (فيشة كهربائية)'), (img7, 'جهاز المكيف (ترمينالات)')]):
        if os.path.exists(img_path):
            c = pics_table2.cell(0, col_idx)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run()
            r.add_picture(img_path, width=Cm(7))
            p2 = c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run(caption)
            set_font(r2, size=10, bold=True, color=(0x1A,0x56,0x9A))

doc.add_paragraph()

add_heading(doc, 'مواصفات جهاز البراد', level=2)
specs_brad = [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال', 'ESP-NOW لاسلكي'),
    ('عنصر التحكم', 'ريلي — 10 أمبير / 240V AC'),
    ('الاستجابة', 'خلال 3 ثوانٍ من أمر التشغيل'),
    ('التغذية', 'DC 12V'),
    ('التوصيل', 'فيشة كهربائية عادية — سهل التركيب'),
]
add_specs_table(doc, '', specs_brad)

add_heading(doc, 'مواصفات جهاز المكيف', level=2)
specs_mkayef = [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال', 'ESP-NOW لاسلكي'),
    ('عنصر التحكم', 'ريلي — 30 أمبير / 240V AC'),
    ('التركيب', 'ترمينالات — للتركيب الثابت'),
    ('مناسب لـ', 'المكيفات، الغسالات، الأحمال الكبيرة'),
    ('التغذية', 'DC 12V'),
]
add_specs_table(doc, '', specs_mkayef)

add_heading(doc, 'مميزات مهمة', level=2)
brad_features = [
    'ريلي داخلي يتحمل حتى 10 أمبير للبراد/الفريزر',
    'نسخة المكيف بريلي 30 أمبير للأحمال الكبيرة',
    'يشتغل عند بلوغ نسبة الشحن المضبوطة',
    'يوقف تلقائياً لحماية البطارية',
    '3 أوقات تأخير إقلاع لحماية الكمبروسر',
    'خيار التشغيل الفوري عند عودة الكهرباء',
]
for f in brad_features:
    add_bullet(doc, f)

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 5: السخان
# ============================================================
add_heading(doc, 'ثالثاً: جهاز السخان', level=1)
add_separator(doc)
add_body(doc, 'يتحكم في سخان الماء بشكل تدريجي (0% إلى 100%) بناءً على الطاقة المتاحة من الألواح الشمسية، مما يعظّم استهلاك الطاقة الشمسية الفائضة.', size=12)

img8 = f"{IMAGES_DIR}/image8.png"
if os.path.exists(img8):
    doc.add_paragraph()
    p_i = doc.add_paragraph()
    p_i.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_i = p_i.add_run()
    run_i.add_picture(img8, width=Cm(9))
    p_c = doc.add_paragraph()
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_c.add_run('جهاز السخان BARQ')
    set_font(r_c, size=10, color=(0x77, 0x77, 0x77))

doc.add_paragraph()

add_heading(doc, 'مواصفات جهاز السخان', level=2)
specs_sakhan = [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال', 'ESP-NOW لاسلكي'),
    ('عنصر التحكم', 'ترياك BTA40'),
    ('نطاق التحكم', '0% إلى 100% بخطوات 1%'),
    ('الحماية', 'إيقاف فوري عند انخفاض الشحن'),
    ('الشاشة', 'رقمية — تعرض الجهد والتيار'),
    ('التوصيل', 'فاز + نتر مدخل، فاز مخرج للسخان'),
]
add_specs_table(doc, '', specs_sakhan)

add_heading(doc, 'ميزة السخان الفريدة', level=2)
add_info_box(doc, '🔥 تحكم تدريجي بالطاقة',
    'على عكس البراد الذي يعمل أو لا يعمل فقط، يمكن للسخان أن يعمل بنسبة جزئية من الطاقة. فعند وجود 2 كيلوواط فائض من الألواح الشمسية، يشغّل السخان بـ 40% فقط — مما يحقق أقصى استفادة من الطاقة الشمسية دون إهدار.')

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 6: دليل التوصيل
# ============================================================
add_heading(doc, 'دليل التوصيل الكهربائي', level=1)
add_separator(doc)

add_heading(doc, 'توصيل المرسلة', level=2)
add_step_box(doc, 1, 'توصيل كابل الشبكة',
    'نوصّل كابل شبكة عادي (RJ45) بين مخرج الإنفرتر (COMM) وجهاز المرسلة. بعد التوصيل مباشرة يبدأ الجهاز بالوميض دلالةً على أنه يعمل.')

add_step_box(doc, 2, 'تغذية الكهرباء',
    'معظم الأنفرترات تغذي المرسلة تلقائياً عبر كابل COMM. وقد يحتاج بعض أنواع الأنفرترات إلى توصيل مدخل DC 12V منفصل.')

add_step_box(doc, 3, 'التحقق من التشغيل',
    'عند التشغيل الصحيح يبدأ الليد بالوميض بشكل منتظم. إذا لم يضيء الليد، تحقق من اتجاه كابل RJ45 أو جرب توصيل DC 12V.')

doc.add_paragraph()

add_heading(doc, 'توصيل المستقبلات', level=2)

add_info_box(doc, '🔌 جهاز البراد',
    'التوصيل بسيط جداً — ضعه في أي مأخذ كهربائي عادي، ثم صِل البراد/الفريزر بالفيشة الخارجة من الجهاز.')

add_info_box(doc, '🏠 جهاز المكيف',
    'له ترمينالات للتركيب الثابت:\n• مدخل فاز (من الكهرباء)\n• مدخل نتر مشترك (تغذية الجهاز)\n• مخرج فاز (موصول بالمكيف)')

add_info_box(doc, '🔥 جهاز السخان',
    'نفس طريقة المكيف بترمينالات:\n• مدخل فاز (من الكهرباء)\n• مدخل نتر مشترك (تغذية الجهاز)\n• مخرج فاز (موصول بالسخان)')

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 7: خطوات الاقتران
# ============================================================
add_heading(doc, 'خطوات الاقتران بين المرسلة والمستقبل', level=1)
add_separator(doc)
add_body(doc, 'بعد توصيل المرسلة والمستقبل (براد أو سخان أو مكيف)، اتبع الخطوات التالية للاقتران:', size=12)
doc.add_paragraph()

add_step_box(doc, 1, 'تفعيل وضع الاقتران في المرسلة',
    'انقر مرتين متتاليتين على زر المرسلة.\nالنتيجة: يضيء الليد الأخضر مرتين ثم يبقى مضيئاً، بينما يبقى الليد الأحمر في وضع الوميض.\nالمعنى: المرسلة جاهزة لعملية الاقتران.')

add_step_box(doc, 2, 'تفعيل وضع الاقتران في المستقبل',
    'انتقل للمستقبل (البراد مثلاً) وانقر زره مرتين متتاليتين.\nالنتيجة:\n• صوت مرة واحدة\n• الليد الأحمر والأخضر يومضان معاً\n• ينطفئان معاً\n• صوت ثلاث مرات\n• يعودان للوميض معاً\n• صوت مرة واحدة\n• صوت مرتين = تم الاقتران بنجاح ✓\nالعملية كلها تتم خلال 30 ثانية أو أقل.')

add_step_box(doc, 3, 'التحقق من نجاح الاقتران',
    'بعد الاقتران الناجح، انقر زر المستقبل (البراد) نقرة واحدة.\nالنتيجة: يبدأ الوميض المتعاكس والسريع (أحمر ثم أخضر بسرعة).\nالمعنى: بدأ الجهاز ببث شبكة BARQ عبر WiFi.')

add_step_box(doc, 4, 'الاتصال بشبكة BARQ',
    'افتح إعدادات WiFi في هاتفك وابحث عن شبكة اسمها BARQ.\nكلمة المرور للمرة الأولى: 11111111 (رقم 1 ثماني مرات)\nبعد الاتصال، الجهاز جاهز للضبط عبر التطبيق.')

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 8: الضبط عبر التطبيق
# ============================================================
add_heading(doc, 'الضبط عبر تطبيق BARQ', level=1)
add_separator(doc)
add_body(doc, 'بعد الاتصال بشبكة BARQ، افتح التطبيق وستجد أمامك خيارين:', size=12)
doc.add_paragraph()

add_info_box(doc, '📱 الضبط المحلي',
    'للضبط المباشر عبر شبكة WiFi الخاصة بالجهاز. لا يحتاج إنترنت. مناسب للضبط الأولي أثناء التركيب.')

add_info_box(doc, '🌐 ضبط الإنترنت',
    'للضبط عن بُعد من أي مكان في العالم. يحتاج أن تكون المرسلة متصلة بالإنترنت.')

doc.add_paragraph()
add_heading(doc, 'حقول الضبط المحلي', level=2)
add_body(doc, 'ادخل الضبط المحلي وأكمل الحقول التالية:', size=12)
doc.add_paragraph()

settings_data = [
    ('الحقل', 'الوصف', 'مثال'),
    ('الاسم', 'اسم تعريفي للجهاز', 'براد المطبخ'),
    ('الهدف', 'ما هو الجهاز الموصول', 'ثلاجة / مكيف / سخان'),
    ('نسبة التشغيل (%)', 'نسبة شحن البطارية التي يبدأ عندها الجهاز بالعمل', '70%'),
    ('نسبة الإيقاف (%)', 'نسبة شحن البطارية التي يتوقف عندها الجهاز', '40%'),
    ('زمن التأخير الأول (ث)', 'تأخير الإقلاع الأول — مفيد عند وجود أكثر من جهاز لتجنب الإقلاع المتزامن', '30 ثانية'),
    ('زمن التأخير الثاني (ث)', 'وقت الانتظار قبل المحاولة الثانية في حال فشل الإقلاع', '60 ثانية'),
    ('زمن التأخير الثالث (ث)', 'وقت الانتظار قبل المحاولة الثالثة في حال فشل الإقلاع مرة ثانية', '120 ثانية'),
    ('زمن تأخير الإيقاف (ث)', 'وقت الانتظار قبل الإيقاف عند انقطاع التيار — يحمي الأجهزة من الانطفاء المفاجئ', '30 ثانية (موصى به)'),
    ('العمل عند قدوم الكهرباء', 'هل يعمل الجهاز فوراً عند عودة التيار الكهربائي؟', 'حسب الرغبة'),
]

settings_table = doc.add_table(rows=len(settings_data), cols=3)
settings_table.style = 'Table Grid'
settings_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
for i, (col1, col2, col3) in enumerate(settings_data):
    row = settings_table.rows[i]
    c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
    if i == 0:
        for c, t in [(c0,col1),(c1,col2),(c2,col3)]:
            set_cell_bg(c, '1A569A')
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(t)
            set_font(r, size=11, bold=True, color=(0xFF,0xFF,0xFF))
    else:
        bg = 'F7FAFF' if i % 2 == 1 else 'FFFFFF'
        for c, t, bld in [(c0,col1,True),(c1,col2,False),(c2,col3,False)]:
            set_cell_bg(c, bg)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)
            r = p.add_run(t)
            set_font(r, size=10, bold=bld, color=(0x22,0x22,0x22))

doc.add_paragraph()
add_info_box(doc, '✅ حفظ الإعدادات',
    'بعد تعبئة جميع الحقول، اضغط زر "حفظ".\nعند الحفظ الناجح: يصدر الجهاز صوتاً مزدوجاً تأكيداً لحفظ المعلومات.')

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 9: إعداد المرسلة للإنترنت
# ============================================================
add_heading(doc, 'ربط المرسلة بالإنترنت (اختياري)', level=1)
add_separator(doc)
add_body(doc, 'في حال أراد العميل المراقبة عن بُعد، يتم ربط المرسلة بالراوتر المنزلي. هذه الخطوة اختيارية — النظام يعمل بالكامل بدونها محلياً.', size=12)
doc.add_paragraph()

add_step_box(doc, 1, 'تفعيل وضع WiFi في المرسلة',
    'انقر زر المرسلة مرة واحدة فقط.\nتبدأ المرسلة ببث شبكة WiFi باسم BARQ.')

add_step_box(doc, 2, 'الاتصال بشبكة BARQ',
    'افتح WiFi في هاتفك واتصل بشبكة BARQ.\nكلمة المرور: 11111111')

add_step_box(doc, 3, 'إدخال بيانات الراوتر المنزلي',
    'افتح التطبيق > اعدادات WiFi.\nأدخل اسم شبكة الراوتر المنزلي (2.4 GHz) وكلمة مروره.\nملاحظة: تأكد أن الراوتر يعمل على تردد 2.4 GHz وليس 5 GHz.')

add_step_box(doc, 4, 'التحقق من الاتصال',
    'بعد الحفظ تعود المرسلة للعمل العادي وتتصل بالراوتر تلقائياً.\nيمكنك الآن متابعة نظامك من أي مكان عبر التطبيق.')

doc.add_paragraph()

add_heading(doc, 'المراقبة عبر التطبيق', level=2)
monitoring_features = [
    ('✅ مع الإنترنت', ['مراقبة من أي مكان في العالم', 'ضبط الإعدادات عن بُعد', 'مخططات بيانية محفوظة', 'إشعارات فورية على الهاتف']),
    ('🔒 بدون إنترنت', ['التحكم الكامل بالأجهزة محلياً', 'آمن وخاص — يعمل 24/7', 'الضبط عبر هاتفك مباشرة', 'لا حاجة لاشتراك أو خدمة خارجية']),
]
for title, items in monitoring_features:
    add_heading(doc, title, level=3)
    for item in items:
        add_bullet(doc, item)
    doc.add_paragraph()

p_br = doc.add_paragraph()
p_br.add_run().add_break(WD_BREAK.PAGE)

# ============================================================
# PAGE 10: المواصفات الفنية الكاملة
# ============================================================
add_heading(doc, 'المواصفات الفنية الكاملة', level=1)
add_separator(doc)

add_specs_table(doc, '📡 المرسلة', [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال بالإنفرتر', 'RS-232 / RJ45 (مخرج COMM)'),
    ('الاتصال اللاسلكي', 'ESP-NOW — مدى يصل 200 متر'),
    ('الاتصال بالإنترنت', 'WiFi 2.4GHz'),
    ('التغذية', 'DC 12V'),
    ('دورة رفع البيانات', 'كل دقيقتين تلقائياً'),
])

add_specs_table(doc, '❄️ جهاز البراد', [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال', 'ESP-NOW لاسلكي'),
    ('عنصر التحكم', 'ريلي — 10 أمبير / 240V AC'),
    ('الاستجابة', 'خلال 3 ثوانٍ من أمر التشغيل'),
    ('التغذية', 'DC 12V'),
])

add_specs_table(doc, '🏠 جهاز المكيف', [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال', 'ESP-NOW لاسلكي'),
    ('عنصر التحكم', 'ريلي — 30 أمبير / 240V AC'),
    ('التركيب', 'ترمينالات — للتركيب الثابت'),
    ('مناسب لـ', 'المكيفات، الغسالات، الأحمال الكبيرة'),
])

add_specs_table(doc, '🔥 جهاز السخان', [
    ('المتحكم', 'ESP8266 / ESP-12E'),
    ('الاتصال', 'ESP-NOW لاسلكي'),
    ('عنصر التحكم', 'ترياك BTA40'),
    ('نطاق التحكم', '0% إلى 100% بخطوات 1%'),
    ('الحماية', 'إيقاف فوري عند انخفاض الشحن'),
    ('الشاشة', 'رقمية — تعرض الجهد والتيار'),
])

add_info_box(doc, '🇸🇾 صناعة سورية',
    'جميع أجهزة نظام برق مصنوعة ومطورة محلياً في سوريا.\nالدعم الفني متوفر مباشرة من الفريق المصنّع.\nجميع الحقوق محفوظة لشركة برق لحلول الطاقة النظيفة.')

# Save
output_path = '/home/user/cataloc-for-barq-232/كتالوج_برق_BARQ.docx'
doc.save(output_path)
print(f"Saved to: {output_path}")
